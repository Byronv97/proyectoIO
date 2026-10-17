"""Build an editable Word version of the academic project report."""

from pathlib import Path
import re

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).parent
SOURCE = ROOT / "PROJECT_DRAFT.md"
OUTPUT = ROOT / "INFORME_PROYECTO_EDITABLE.docx"
GRAPH_PNG = ROOT / "assets" / "region_factible.png"


def clean_inline(text: str) -> str:
    text = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"__([^_]+)__", r"\1", text)
    text = text.replace("`", "")
    return text


def equation_to_text(equation: str) -> str:
    equation = equation.strip()
    equation = re.sub(r"\\text\{([^}]*)\}", r"\1", equation)
    equation = equation.replace(r"\leq", "<=")
    equation = equation.replace(r"\geq", ">=")
    equation = equation.replace(r"\qquad", "   ")
    equation = equation.replace(r"\quad", "   ")
    equation = equation.replace("{", "").replace("}", "")
    return equation


def create_graph_png() -> None:
    figure, axis = plt.subplots(figsize=(8, 5.2))
    x_values = [index * 60 / 200 for index in range(201)]
    constraints = (
        ("Harina: 2x + 3y <= 120", 2, 3, 120, "#3b82f6"),
        ("Mano de obra: 3x + 2y <= 120", 3, 2, 120, "#ef4444"),
        ("Capacidad: x + y <= 45", 1, 1, 45, "#8b5cf6"),
    )
    for label, coefficient_x, coefficient_y, available, color in constraints:
        axis.plot(
            x_values,
            [(available - coefficient_x * x) / coefficient_y for x in x_values],
            color=color,
            label=label,
        )
    polygon = [(0, 0), (40, 0), (30, 15), (15, 30), (0, 40)]
    axis.fill(
        [point[0] for point in polygon],
        [point[1] for point in polygon],
        color="#2f7f75",
        alpha=0.2,
        label="Region factible",
    )
    axis.scatter([15], [30], color="#e07a5f", marker="*", s=220, label="Optimo (15, 30)")
    axis.set_xlim(0, 65)
    axis.set_ylim(0, 70)
    axis.set_xlabel("Lotes de galletas (x)")
    axis.set_ylabel("Lotes de cupcakes (y)")
    axis.set_title("Region factible y solucion optima")
    axis.grid(alpha=0.25)
    axis.legend(fontsize=8, loc="upper right")
    figure.tight_layout()
    figure.savefig(GRAPH_PNG, dpi=180, bbox_inches="tight")
    plt.close(figure)


def set_document_style(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    for style_name, size, color in (
        ("Title", 22, RGBColor(18, 52, 59)),
        ("Heading 1", 16, RGBColor(18, 52, 59)),
        ("Heading 2", 12.5, RGBColor(47, 127, 117)),
    ):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = color


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, rows[0]):
        cell.text = clean_inline(value)
    for row in rows[1:]:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = clean_inline(value)
    document.add_paragraph()


def add_paragraph(document: Document, text: str, style: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    paragraph.add_run(clean_inline(text))


def build_document() -> None:
    document = Document()
    set_document_style(document)
    document.core_properties.title = "Aplicación de la programación lineal para optimizar la producción de Dulce Trigo"
    document.core_properties.subject = "Proyecto final de Investigación de Operaciones"

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            text = clean_inline(heading.group(2))
            if level == 1:
                paragraph = document.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run(text)
            else:
                document.add_heading(text, level=level - 1)
            index += 1
            continue

        if line.startswith("> "):
            quote_lines = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip().lstrip("> "))
                index += 1
            paragraph = document.add_paragraph(style="Intense Quote")
            paragraph.add_run(clean_inline(" ".join(quote_lines)))
            continue

        image = re.match(r"!\[[^]]*\]\(([^)]+)\)", line)
        if image:
            image_path = ROOT / image.group(1)
            if image_path.suffix.lower() == ".svg":
                create_graph_png()
                image_path = GRAPH_PNG
            if image_path.exists():
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(str(image_path), width=Inches(6.2))
            index += 1
            continue

        if line == r"\[":
            equation_lines = []
            index += 1
            while index < len(lines) and lines[index].strip() != r"\]":
                equation_lines.append(lines[index])
                index += 1
            index += 1
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run(equation_to_text(" ".join(equation_lines))).font.name = "Courier New"
            continue

        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                values = [value.strip() for value in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", value) for value in values):
                    rows.append(values)
                index += 1
            add_table(document, rows)
            continue

        if re.match(r"^[-*]\s+", line):
            add_paragraph(document, re.sub(r"^[-*]\s+", "", line), "List Bullet")
            index += 1
            continue

        if re.match(r"^\d+\.\s+", line):
            add_paragraph(document, re.sub(r"^\d+\.\s+", "", line), "List Number")
            index += 1
            continue

        paragraph_lines = [line]
        index += 1
        while index < len(lines):
            next_line = lines[index].strip()
            if (
                not next_line
                or re.match(r"^(#{1,3})\s+", next_line)
                or next_line.startswith(("> ", "|", "!", "\\[", "- ", "* "))
                or re.match(r"^\d+\.\s+", next_line)
            ):
                break
            paragraph_lines.append(next_line)
            index += 1
        add_paragraph(document, " ".join(paragraph_lines))

    document.save(OUTPUT)
    print(f"Word generado: {OUTPUT}")


if __name__ == "__main__":
    build_document()
